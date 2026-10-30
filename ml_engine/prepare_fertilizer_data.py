from pathlib import Path
import sys
ROOT = Path(__file__).resolve().parent.parent

# Tell Python where the project root is
sys.path.insert(0, str(ROOT))

from ml_engine.fertilizer_units import convert_to_kg
from pathlib import Path
import csv
from collections import Counter
from openpyxl import load_workbook

from ml_engine.fertilizer_units import convert_to_kg

ROOT = Path(__file__).resolve().parent.parent

input_file = ROOT / "Carbon Farming HH 2023 CIMMYT Survey Data_1.xlsx"
output_file = ROOT / "outputs" / "fertilizer_cleaned.csv"

# Open the Excel workbook
wb = load_workbook(input_file, read_only=True, data_only=True)

try:
    sheet_name = "Chem_fert_kh_details"

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}")

    ws = wb[sheet_name]
    headers = [cell.value for cell in ws[1]]

    quantity_column = "Qty_inp_season_chem_fert_kh"
    unit_column = "Chem_fert_inp_unit_kh"

    # Check that required columns exist
    missing_columns = [
        col for col in (quantity_column, unit_column)
        if col not in headers
    ]

    if missing_columns:
        raise ValueError(f"Missing columns: {missing_columns}")

    # Create the output folder if it does not exist
    output_file.parent.mkdir(parents=True, exist_ok=True)

    output_headers = headers + [
        "quantity_kg",
        "conversion_status"
    ]

    counts = Counter()
    total = 0

    with output_file.open(
        "w", newline="", encoding="utf-8-sig"
    ) as file:

        writer = csv.DictWriter(file, fieldnames=output_headers)
        writer.writeheader()

        for values in ws.iter_rows(min_row=2, values_only=True):

            # Skip completely empty rows
            if all(
                value is None or
                (isinstance(value, str) and not value.strip())
                for value in values
            ):
                continue

            record = dict(zip(headers, values))

            quantity = record[quantity_column]
            unit = record[unit_column]

            quantity_missing = (
                quantity is None or
                (isinstance(quantity, str) and not quantity.strip())
            )
            unit_missing = (
                unit is None or
                (isinstance(unit, str) and not unit.strip())
            )

            # Missing information is NOT automatically zero
            if quantity_missing and unit_missing:
                kg, status = None, "missing_quantity_and_unit"
            elif quantity_missing:
                kg, status = None, "missing_quantity"
            elif unit_missing:
                kg, status = None, "missing_unit"
            else:
                kg, status = convert_to_kg(quantity, unit)

            record["quantity_kg"] = kg
            record["conversion_status"] = status

            writer.writerow(record)

            counts[status] += 1
            total += 1

finally:
    wb.close()

print("\nFertilizer preprocessing complete")
print("Total records:", total)
print("Conversion results:")

for status, count in sorted(counts.items()):
    print(f"  {status}: {count}")

print("Output file:", output_file)
